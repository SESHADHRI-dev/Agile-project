import csv
import io
from fastapi import APIRouter, Depends, Query, Response, HTTPException
from typing import Dict, Any, Optional
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
from backend.app.database import db
from backend.app.auth import get_current_user
from backend.app.config import STORAGE_MODE, S3_REPORTS_BUCKET, AWS_REGION, REPORTS_DIR
from ml.forecaster import DemandForecaster

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

router = APIRouter(prefix="/reports", tags=["Reports"])

IST = ZoneInfo("Asia/Kolkata")


def _get_ist_timestamp_str() -> str:
    return datetime.now(IST).strftime("%d/%m/%Y, %I:%M:%S %p IST")


def _get_file_timestamp_str() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")


@router.get("/summary", response_model=Dict[str, Any])
def get_reports_summary(user: dict = Depends(get_current_user)):
    """Executive reporting KPIs."""
    products = db.get_products()
    purchases = db.get_purchases(limit=100)
    sales = db.get_sales(limit=100)
    
    total_sales_revenue = round(sum(s["total_revenue"] for s in sales), 2)
    total_purchase_spend = round(sum(p["total_cost"] for p in purchases), 2)
    total_inventory_value = round(sum(p["quantity"] * p["price"] for p in products), 2)

    return {
        "success": True,
        "total_products": len(products),
        "total_inventory_value": total_inventory_value,
        "total_sales_revenue": total_sales_revenue,
        "total_purchase_spend": total_purchase_spend,
        "recent_sales_count": len(sales),
        "recent_purchases_count": len(purchases)
    }


def _gather_report_data(report_type: str):
    """Gathers structured tabular data and metadata for each report type."""
    if report_type == "inventory":
        title = "Current Inventory Valuation & Stock Audit"
        headers = ["Product ID", "Product Name", "Category", "HSN Code", "GST Rate (%)", "Unit Price (₹)", "Quantity in Stock", "Min Stock Level", "Status", "Valuation (₹)"]
        rows = []
        for p in db.get_products():
            val = round(p["quantity"] * p["price"], 2)
            rows.append([
                p["id"],
                p["name"],
                p["category"],
                p.get("hsn_code", "8536"),
                float(p.get("gst_rate", 18.0)),
                float(p["price"]),
                int(p["quantity"]),
                int(p["min_stock_level"]),
                p["status"],
                val
            ])
        return title, headers, rows

    elif report_type == "low_stock":
        title = "Low-Stock & Critical Deficit Audit"
        headers = ["Product ID", "Product Name", "Category", "Current Stock", "Min Stock Level", "Deficit Units", "Status"]
        rows = []
        for p in db.get_products():
            if p["status"] in ["LOW STOCK", "OUT OF STOCK"]:
                deficit = max(0, p["min_stock_level"] - p["quantity"])
                rows.append([
                    p["id"],
                    p["name"],
                    p["category"],
                    int(p["quantity"]),
                    int(p["min_stock_level"]),
                    int(deficit),
                    p["status"]
                ])
        return title, headers, rows

    elif report_type == "sales":
        title = "Customer Outbound Sales Transaction Ledger"
        headers = ["Sale ID", "Product ID", "Product Name", "Customer Name", "Quantity Sold", "Unit Price (₹)", "Total Revenue (₹)", "GST (%)", "CGST (₹)", "SGST (₹)", "IGST (₹)", "Sale Date (IST)", "Recorded By"]
        rows = []
        for s in db.get_sales(limit=1000):
            rows.append([
                s["id"],
                s["product_id"],
                s.get("product_name", ""),
                s.get("customer_name", "Sri Ganesh Traders"),
                int(s["quantity"]),
                float(s["unit_price"]),
                float(s["total_revenue"]),
                float(s.get("gst_rate", 18.0)),
                float(s.get("cgst", 0.0)),
                float(s.get("sgst", 0.0)),
                float(s.get("igst", 0.0)),
                s["sale_date"],
                s["created_by"]
            ])
        return title, headers, rows

    elif report_type == "purchases":
        title = "Supplier Inbound Procurement Ledger"
        headers = ["Purchase ID", "Product ID", "Product Name", "Supplier Name", "Supplier GSTIN", "Quantity", "Unit Cost (₹)", "Total Cost (₹)", "GST (%)", "CGST (₹)", "SGST (₹)", "Purchase Date (IST)", "Recorded By"]
        rows = []
        for p in db.get_purchases(limit=1000):
            rows.append([
                p["id"],
                p["product_id"],
                p.get("product_name", ""),
                p.get("supplier_name", ""),
                p.get("supplier_gstin", "33AABCS1234A1Z1"),
                int(p["quantity"]),
                float(p["unit_cost"]),
                float(p["total_cost"]),
                float(p.get("gst_rate", 18.0)),
                float(p.get("cgst", 0.0)),
                float(p.get("sgst", 0.0)),
                p["purchase_date"],
                p["created_by"]
            ])
        return title, headers, rows

    elif report_type == "predictions":
        title = "Demand Forecasting & Restocking Recommendations"
        headers = ["Product ID", "Product Name", "Category", "Current Stock", "Forecast Period (Days)", "Predicted Demand", "Safety Stock", "Recommended Restock", "Urgency Status"]
        rows = []
        for prod in db.get_products():
            sales = db.get_sales(limit=200, product_id=prod["id"])
            rec = DemandForecaster.generate_recommendation(sales_records=sales, current_stock=prod["quantity"])
            rows.append([
                prod["id"],
                prod["name"],
                prod["category"],
                int(prod["quantity"]),
                int(rec["forecast_horizon_days"]),
                int(rec["predicted_demand"]),
                int(rec["safety_stock"]),
                int(rec["recommended_restock"]),
                rec["urgency_status"]
            ])
        return title, headers, rows

    raise HTTPException(status_code=400, detail=f"Unsupported report type '{report_type}'")


def _generate_excel_workbook(title: str, headers: list, rows: list, user_email: str) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:30]

    # Styles
    title_font = Font(name="Segoe UI", size=15, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Segoe UI", size=9, italic=True, color="475569")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    regular_font = Font(name="Segoe UI", size=9, color="0F172A")
    border_thin = Side(border_style="thin", color="E2E8F0")
    cell_border = Border(top=border_thin, left=border_thin, right=border_thin, bottom=border_thin)

    # Title Banner
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(row=1, column=1, value=f"IntelliStock India — {title}").font = title_font

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    gen_text = f"Hub: Katpadi, Vellore, Tamil Nadu | Generated: {_get_ist_timestamp_str()} | Authorized By: {user_email} | Currency: INR (₹)"
    ws.cell(row=2, column=1, value=gen_text).font = subtitle_font

    # Header Row
    header_row_idx = 4
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border

    # Data Rows
    for row_idx, r in enumerate(rows, start=header_row_idx + 1):
        is_alt = (row_idx % 2 == 0)
        for col_idx, val in enumerate(r, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = regular_font
            cell.border = cell_border
            if is_alt:
                cell.fill = alt_fill
            
            # Numeric formatting
            if isinstance(val, float):
                cell.number_format = "₹#,##0.00"
                cell.alignment = Alignment(horizontal="right")
            elif isinstance(val, int):
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="left")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len and cell.row >= header_row_idx:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _generate_pdf_document(title: str, headers: list, rows: list, user_email: str) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(letter),
        leftMargin=20,
        rightMargin=20,
        topMargin=20,
        bottomMargin=20
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#1E3A8A")
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#475569")
    )
    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7.5,
        leading=9,
        textColor=colors.HexColor("#0F172A")
    )
    header_cell_style = ParagraphStyle(
        "TableHeaderCell",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1  # Center
    )

    elements = [
        Paragraph(f"<b>IntelliStock India</b> &mdash; {title}", title_style),
        Paragraph(
            f"Katpadi, Vellore, Tamil Nadu &bull; Generated: {_get_ist_timestamp_str()} &bull; Auditor: {user_email} &bull; Currency: INR (&inr;)",
            subtitle_style
        ),
        Spacer(1, 10)
    ]

    # Convert all cell values into wrapped Paragraphs or strings
    table_data = []
    # Header row
    table_data.append([Paragraph(h, header_cell_style) for h in headers])
    for r in rows:
        row_cells = []
        for val in r:
            if isinstance(val, float):
                formatted = f"₹{val:,.2f}"
            elif isinstance(val, int):
                formatted = f"{val:,}"
            else:
                formatted = str(val)
            row_cells.append(Paragraph(formatted, cell_style))
        table_data.append(row_cells)

    t = Table(table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(t)
    doc.build(elements)
    return buf.getvalue()


@router.get("/export")
def export_report(
    report_type: str = Query("inventory", pattern="^(inventory|sales|purchases|predictions|low_stock)$"),
    format: str = Query("csv", pattern="^(csv|xlsx|excel|pdf)$"),
    user: dict = Depends(get_current_user)
):
    """
    Exports an enterprise inventory audit report in CSV, Excel (.xlsx), or PDF format.
    Includes Indian localization (INR ₹, GST breakdowns, Asia/Kolkata timestamps).
    """
    title, headers, rows = _gather_report_data(report_type)
    file_ts = _get_file_timestamp_str()
    user_email = user.get("username", "admin@intellistock.in")

    # Format 1: Excel (.xlsx)
    if format in ["xlsx", "excel"]:
        filename = f"report_{report_type}_{file_ts}.xlsx"
        content_bytes = _generate_excel_workbook(title, headers, rows, user_email)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    # Format 2: PDF (.pdf)
    elif format == "pdf":
        filename = f"report_{report_type}_{file_ts}.pdf"
        content_bytes = _generate_pdf_document(title, headers, rows, user_email)
        media_type = "application/pdf"

    # Format 3: CSV (.csv)
    else:
        filename = f"report_{report_type}_{file_ts}.csv"
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        for r in rows:
            formatted_row = []
            for val in r:
                if isinstance(val, float):
                    formatted_row.append(f"₹{val:.2f}")
                else:
                    formatted_row.append(val)
            writer.writerow(formatted_row)
        content_bytes = output.getvalue().encode("utf-8")
        media_type = "text/csv; charset=utf-8"

    # Save artifact copy locally in REPORTS_DIR
    try:
        local_path = REPORTS_DIR / filename
        if isinstance(content_bytes, bytes):
            local_path.write_bytes(content_bytes)
        else:
            local_path.write_text(content_bytes, encoding="utf-8")
    except Exception as e:
        pass

    # Optional S3 upload if in live AWS mode
    s3_url = None
    if STORAGE_MODE == "aws":
        try:
            import boto3
            s3 = boto3.client("s3", region_name=AWS_REGION)
            s3.put_object(
                Bucket=S3_REPORTS_BUCKET,
                Key=f"reports/{filename}",
                Body=content_bytes,
                ContentType=media_type
            )
            s3_url = f"https://{S3_REPORTS_BUCKET}.s3.{AWS_REGION}.amazonaws.com/reports/{filename}"
        except Exception:
            pass

    return Response(
        content=content_bytes,
        media_type=media_type,
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "X-Report-Filename": filename,
            "X-Report-Format": format,
            "X-S3-URL": s3_url or ""
        }
    )
